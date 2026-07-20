# -*- coding: utf-8 -*-
from __future__ import annotations

import glob
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook


@dataclass
class KnowledgeSheet:
    name: str
    rows: List[Dict[str, Any]] = field(default_factory=list)


@dataclass
class ItemKnowledge:
    questionIndex: int
    file_path: str
    sheets: Dict[str, KnowledgeSheet]

    def sheet_rows(self, sheet_name: str) -> List[Dict[str, Any]]:
        if sheet_name not in self.sheets:
            return []
        return self.sheets[sheet_name].rows

    def ai_rules(self) -> List[Dict[str, Any]]:
        return self.sheet_rows("15AI访谈规则") or self.sheet_rows("15 AI访谈规则")

    def evidence_schema_rows(self) -> List[Dict[str, Any]]:
        return self.sheet_rows("16访谈输出证据规范") or self.sheet_rows("16 访谈输出证据规范")

    def relevant_snippets(self, task_card: Dict[str, Any], max_rows_per_sheet: int = 8) -> str:
        """为模型拼接可控长度的知识库片段。

        这里不一次性塞入完整Excel，而是选取访谈中最常用的表：选项解释、访谈触发、追问脚本、评分锚点、AI访谈规则。
        """
        important_sheets = ["05选项解释", "08访谈触发", "09追问脚本", "10评分锚点", "15AI访谈规则", "16访谈输出证据规范"]
        parts: List[str] = []
        for name in important_sheets:
            rows = self.sheet_rows(name)
            if not rows:
                continue
            parts.append(f"【{name}】")
            for row in rows[:max_rows_per_sheet]:
                cleaned = {str(k): v for k, v in row.items() if v not in (None, "")}
                if cleaned:
                    parts.append(json.dumps(cleaned, ensure_ascii=False))
        return "\n".join(parts)


class KnowledgeLoader:
    def __init__(self, knowledge_dir: str | Path) -> None:
        self.knowledge_dir = Path(knowledge_dir)
        self.items: Dict[int, ItemKnowledge] = {}

    def load_all(self) -> Dict[int, ItemKnowledge]:
        if not self.knowledge_dir.exists():
            raise FileNotFoundError(f"知识库目录不存在：{self.knowledge_dir}")
        files = sorted(glob.glob(str(self.knowledge_dir / "第*题_*知识库*.xlsx")))
        if not files:
            files = sorted(glob.glob(str(self.knowledge_dir / "*.xlsx")))
        for file_path in files:
            q = self._infer_question_index(Path(file_path).name)
            if q is None or not (1 <= q <= 10):
                continue
            self.items[q] = self._load_item(q, file_path)
        missing = [q for q in range(1, 11) if q not in self.items]
        if missing:
            raise ValueError(f"知识库缺少以下题号文件：{missing}；当前目录：{self.knowledge_dir}")
        return self.items

    def get(self, question_index: int) -> ItemKnowledge:
        if not self.items:
            self.load_all()
        return self.items[int(question_index)]

    def _infer_question_index(self, filename: str) -> Optional[int]:
        import re
        m = re.search(r"第\s*(\d+)\s*题", filename)
        return int(m.group(1)) if m else None

    def _load_item(self, question_index: int, file_path: str) -> ItemKnowledge:
        wb = load_workbook(file_path, data_only=True, read_only=True)
        sheets: Dict[str, KnowledgeSheet] = {}
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                sheets[ws.title] = KnowledgeSheet(ws.title, [])
                continue
            headers = [str(x).strip() if x is not None else f"col_{i+1}" for i, x in enumerate(rows[0])]
            data_rows: List[Dict[str, Any]] = []
            for values in rows[1:]:
                row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
                # 跳过全空行。
                if any(v not in (None, "") for v in row.values()):
                    data_rows.append(row)
            sheets[ws.title] = KnowledgeSheet(ws.title, data_rows)
        return ItemKnowledge(question_index, file_path, sheets)


def rules_to_task_fields(rules: List[Dict[str, Any]], task_card: Dict[str, Any]) -> Dict[str, List[str] | str]:
    """根据15AI访谈规则生成任务卡中的规则字段。

    兼容不同表头。优先读取 rule_type/rule_content/task_card_target_field；若表头不一致，尽量猜测。
    """
    teacher_profile = task_card.get("teacher_answer_profile", {})
    final_order = teacher_profile.get("teacherFinalOrder", "")
    priority_option = teacher_profile.get("priorityOption", "")
    priority_pair = teacher_profile.get("priorityPair", "")

    out: Dict[str, List[str] | str] = {
        "interview_hypotheses": [],
        "must_obtain_evidence": [],
        "recommended_probes": [],
        "recommended_interview_flow": [],
    }

    def match_condition(cond: str) -> bool:
        cond = (cond or "all").strip()
        if cond.lower() in {"all", "全部", "通用", ""}:
            return True
        cond_upper = cond.upper().replace(" ", "")
        # 简易触发条件。
        if "A_FIRST" in cond_upper or "A首" in cond:
            return final_order.startswith("A")
        if "B_FIRST" in cond_upper or "B首" in cond:
            return final_order.startswith("B")
        if "C_FIRST" in cond_upper or "C首" in cond:
            return final_order.startswith("C")
        if "D_FIRST" in cond_upper or "D首" in cond:
            return final_order.startswith("D")
        if "PRIORITYOPTION" in cond_upper:
            return priority_option and priority_option in cond_upper
        if "PRIORITYPAIR" in cond_upper:
            return priority_pair and priority_pair.replace("/", "") in cond_upper.replace("/", "")
        if priority_option and priority_option in cond_upper:
            return True
        if priority_pair and priority_pair.replace("/", "") in cond_upper.replace("/", ""):
            return True
        return False

    for row in sorted(rules, key=lambda r: str(r.get("priority", r.get("优先级", "99")))):
        rule_type = str(row.get("rule_type") or row.get("规则类型") or row.get("类型") or "").strip()
        target = str(row.get("task_card_target_field") or row.get("目标字段") or "").strip()
        content = str(row.get("rule_content") or row.get("规则内容") or row.get("content") or row.get("内容") or "").strip()
        cond = str(row.get("trigger_condition") or row.get("触发条件") or "all")
        if not content or not match_condition(cond):
            continue
        low = f"{rule_type} {target}".lower()
        if "main_focus" in low or "主目标" in rule_type:
            out["interview_main_focus"] = content
        elif "secondary" in low or "次目标" in rule_type:
            out["interview_secondary_focus"] = content
        elif "hypothesis" in low or "假设" in rule_type:
            out["interview_hypotheses"].append(content)  # type: ignore[index]
        elif "evidence" in low or "证据" in rule_type:
            out["must_obtain_evidence"].append(content)  # type: ignore[index]
        elif "probe" in low or "追问" in rule_type or "question" in low:
            out["recommended_probes"].append(content)  # type: ignore[index]
        elif "flow" in low or "流程" in rule_type:
            out["recommended_interview_flow"].append(content)  # type: ignore[index]

    return out

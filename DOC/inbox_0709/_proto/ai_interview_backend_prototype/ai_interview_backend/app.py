# -*- coding: utf-8 -*-
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook

from .config import settings
from .schemas import (
    EvidenceTable,
    FinishItemRequest,
    HealthResponse,
    NextMessageRequest,
    NextMessageResponse,
    PrepareMaterialsRequest,
    StartInterviewRequest,
    StartInterviewResponse,
)
from .task_card_builder import prepare_materials
from .interview_state import InterviewService

app = FastAPI(title="幼儿园教师游戏支持与引导AI访谈后端原型", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

_service_cache: Dict[str, InterviewService] = {}
_task_cards_cache: Dict[str, List[Dict[str, Any]]] = {}


def get_service(knowledge_dir: str | None = None) -> InterviewService:
    kd = str(knowledge_dir or settings.knowledge_dir)
    if kd not in _service_cache:
        _service_cache[kd] = InterviewService(kd)
    return _service_cache[kd]


def load_task_cards_from_xlsx(materials_path: str | Path) -> List[Dict[str, Any]]:
    materials_path = str(materials_path)
    if materials_path in _task_cards_cache:
        return _task_cards_cache[materials_path]
    wb = load_workbook(materials_path, data_only=True, read_only=True)
    if "AI后台任务卡" not in wb.sheetnames:
        raise ValueError(f"{materials_path} 中不存在工作表：AI后台任务卡")
    ws = wb["AI后台任务卡"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(x).strip() if x is not None else f"col_{i+1}" for i, x in enumerate(rows[0])]
    data: List[Dict[str, Any]] = []
    for values in rows[1:]:
        row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        if row.get("task_card_json"):
            data.append(row)
    _task_cards_cache[materials_path] = data
    return data


@app.get("/health", response_model=HealthResponse)
async def health() -> HealthResponse:
    return HealthResponse(status="ok", mock_model=settings.mock_model, model_name=settings.model_name)


@app.post("/admin/prepare-materials")
async def api_prepare_materials(req: PrepareMaterialsRequest) -> Dict[str, Any]:
    try:
        return prepare_materials(req.results_csv_path, req.logs_csv_path, req.knowledge_dir, req.output_dir)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/interview/start", response_model=StartInterviewResponse)
async def start_interview(req: StartInterviewRequest) -> StartInterviewResponse:
    try:
        task_json = req.task_card_json
        if not task_json:
            materials_path = Path(settings.materials_dir) / "AI访谈材料.xlsx"
            cards = load_task_cards_from_xlsx(materials_path)
            matched = [r for r in cards if str(r.get("userOpenid")) == req.userOpenid and int(r.get("final_rank") or 0) == req.final_rank]
            if not matched:
                raise ValueError(f"未找到任务卡：userOpenid={req.userOpenid}, final_rank={req.final_rank}。可直接传入task_card_json。")
            task_json = matched[0]["task_card_json"]
        task_card = json.loads(task_json)
        service = get_service(settings.knowledge_dir)
        session = await service.start(task_json)
        return StartInterviewResponse(
            session_id=session.session_id,
            userOpenid=session.userOpenid,
            final_rank=session.final_rank,
            questionIndex=session.questionIndex,
            item_title=session.item_title,
            assistant_message=session.dialogue[-1]["content"],
            stage=session.stage,
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/interview/message", response_model=NextMessageResponse)
async def next_message(req: NextMessageRequest) -> NextMessageResponse:
    try:
        service = get_service(settings.knowledge_dir)
        session = await service.next_question(req.session_id, req.teacher_message)
        return NextMessageResponse(
            session_id=session.session_id,
            assistant_message=session.dialogue[-1]["content"],
            stage=session.stage,
            turn_count=session.turn_count,
            should_finish=session.should_finish(),
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/interview/finish-item", response_model=EvidenceTable)
async def finish_item(req: FinishItemRequest) -> EvidenceTable:
    try:
        service = get_service(settings.knowledge_dir)
        session = await service.finish_item(req.session_id)
        return EvidenceTable(
            session_id=session.session_id,
            userOpenid=session.userOpenid,
            questionIndex=session.questionIndex,
            item_title=session.item_title,
            evidence=session.evidence or {},
            dialogue=session.dialogue,
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
